# Builds docs/SRS.docx from docs/SRS.md.
#
# The markdown file is the source I edit. This script turns it into the Word
# document that gets handed in. Run it from the project root:
#
#     python docs/build-srs.py
#
# It needs python-docx (pip install python-docx). It only understands the bits
# of markdown the SRS actually uses: headings, paragraphs, bullet and numbered
# lists, tables, bold, inline code and horizontal rules.

import re
import os
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, 'SRS.md')
OUT = os.path.join(HERE, 'SRS.docx')

ACCENT = RGBColor(0xC0, 0x1F, 0x2E)     # the site's red
GREY = RGBColor(0x55, 0x5B, 0x63)
HEADER_FILL = 'E9ECEF'


# ---------------------------------------------------------------- helpers
def set_cell_fill(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_fill)
    tc_pr.append(shd)


def add_field(paragraph, instr):
    """Adds a Word field (TOC, PAGE) to a paragraph."""
    run = paragraph.add_run()
    fld_begin = OxmlElement('w:fldChar')
    fld_begin.set(qn('w:fldCharType'), 'begin')
    instr_el = OxmlElement('w:instrText')
    instr_el.set(qn('xml:space'), 'preserve')
    instr_el.text = instr
    fld_sep = OxmlElement('w:fldChar')
    fld_sep.set(qn('w:fldCharType'), 'separate')
    placeholder = OxmlElement('w:t')
    placeholder.text = 'Right-click and choose Update Field to build the table of contents.'
    fld_end = OxmlElement('w:fldChar')
    fld_end.set(qn('w:fldCharType'), 'end')
    for el in (fld_begin, instr_el, fld_sep, placeholder, fld_end):
        run._r.append(el)


def ask_word_to_update_fields(doc):
    """Makes Word refresh the table of contents when the file is opened."""
    settings = doc.settings.element
    update = OxmlElement('w:updateFields')
    update.set(qn('w:val'), 'true')
    settings.append(update)


INLINE = re.compile(r'(\*\*[^*]+\*\*|`[^`]+`)')


def add_inline(target, text, base_bold=False, size=None):
    """Writes text into a paragraph (or a table cell), handling **bold** and `code`."""
    paragraph = target.paragraphs[0] if hasattr(target, 'paragraphs') else target
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith('**') and piece.endswith('**'):
            run = paragraph.add_run(piece[2:-2])
            run.bold = True
        elif piece.startswith('`') and piece.endswith('`'):
            run = paragraph.add_run(piece[1:-1])
            run.font.name = 'Consolas'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Consolas')
            run.font.color.rgb = RGBColor(0x8B, 0x1A, 0x1A)
        else:
            run = paragraph.add_run(piece)
            if base_bold:
                run.bold = True
        if size:
            run.font.size = size


def split_row(line):
    cells = line.strip().strip('|').split('|')
    return [c.strip() for c in cells]


def add_table(doc, rows):
    header, body = split_row(rows[0]), rows[2:]   # rows[1] is the |---| line
    cols = len(header)
    table = doc.add_table(rows=1, cols=cols)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, text in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = ''
        add_inline(cell.paragraphs[0], text, base_bold=True, size=Pt(9.5))
        set_cell_fill(cell, HEADER_FILL)

    body_cells = []
    for row in body:
        cells = split_row(row)
        cells += [''] * (cols - len(cells))
        body_cells.append(cells)
        row_cells = table.add_row().cells
        for i in range(cols):
            row_cells[i].text = ''
            add_inline(row_cells[i].paragraphs[0], cells[i], size=Pt(9.5))

    # give each column a width that matches how much text it holds, so a long
    # description column does not get squeezed to the same size as an id column
    usable = Cm(17)
    longest = [max([len(header[i])] + [len(r[i]) for r in body_cells]) for i in range(cols)]
    # short columns (ids, one word values) get a fixed width, the rest share what is left
    widths = [Cm(2.3) if n <= 12 else None for n in longest]
    remaining = usable - sum(w for w in widths if w)
    flexible = [min(longest[i], 60) for i in range(cols) if widths[i] is None]
    for i in range(cols):
        if widths[i] is None:
            widths[i] = int(remaining * min(longest[i], 60) / sum(flexible))
    table.autofit = False
    for i in range(cols):
        for row in table.rows:
            row.cells[i].width = widths[i]

    doc.add_paragraph()


# ---------------------------------------------------------------- document setup
def build():
    doc = Document()

    # page size and margins
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    for side in ('left_margin', 'right_margin'):
        setattr(section, side, Cm(2))
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)

    styles = doc.styles
    styles['Normal'].font.name = 'Calibri'
    styles['Normal'].font.size = Pt(10.5)
    for name, size in (('Heading 1', 18), ('Heading 2', 14), ('Heading 3', 12)):
        st = styles[name]
        st.font.name = 'Calibri'
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = ACCENT if name == 'Heading 1' else RGBColor(0x22, 0x26, 0x2B)

    # page number in the footer
    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer_p.add_run('Maslul SRS    ')
    r.font.size = Pt(8)
    r.font.color.rgb = GREY
    add_field(footer_p, 'PAGE')

    lines = open(SRC, encoding='utf-8').read().splitlines()

    # ---- title page from the first two "# " lines and the meta table
    i = 0
    title_lines = []
    while i < len(lines) and lines[i].startswith('# '):
        title_lines.append(lines[i][2:].strip())
        i += 1

    for _ in range(6):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title_lines[0])
    run.font.size = Pt(36)
    run.bold = True
    run.font.color.rgb = ACCENT
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title_lines[1] if len(title_lines) > 1 else '')
    run.font.size = Pt(20)
    run.font.color.rgb = RGBColor(0x22, 0x26, 0x2B)
    doc.add_paragraph()

    # meta table straight after the title
    while i < len(lines) and not lines[i].startswith('|'):
        i += 1
    meta = []
    while i < len(lines) and lines[i].startswith('|'):
        meta.append(lines[i])
        i += 1
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in meta[2:]:
        k, v = split_row(row)[:2]
        cells = table.add_row().cells
        add_inline(cells[0], k, base_bold=True)
        add_inline(cells[1], v)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---- table of contents
    h = doc.add_paragraph()
    r = h.add_run('Contents')
    r.bold = True
    r.font.size = Pt(16)
    r.font.color.rgb = ACCENT
    add_field(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u')
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---- the body
    in_table = []

    def flush_table():
        nonlocal in_table
        if in_table:
            add_table(doc, in_table)
            in_table = []

    while i < len(lines):
        line = lines[i]
        i += 1
        stripped = line.strip()

        if stripped.startswith('|'):
            in_table.append(stripped)
            continue
        flush_table()

        if not stripped:
            continue
        if stripped == '---':
            continue

        if stripped.startswith('#### '):
            doc.add_heading(stripped[5:], level=3)
        elif stripped.startswith('### '):
            doc.add_heading(stripped[4:], level=2)
        elif stripped.startswith('## '):
            # every top level chapter starts on a new page
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            doc.add_heading(stripped[3:], level=1)
        elif stripped.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            add_inline(p, stripped[2:])
        elif re.match(r'^\d+\. ', stripped):
            p = doc.add_paragraph(style='List Number')
            add_inline(p, re.sub(r'^\d+\. ', '', stripped))
        else:
            p = doc.add_paragraph()
            add_inline(p, stripped)

    flush_table()
    ask_word_to_update_fields(doc)
    doc.save(OUT)
    print('wrote', OUT)


if __name__ == '__main__':
    build()
